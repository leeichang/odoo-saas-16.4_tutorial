# -*- coding: utf-8 -*-

import base64
import io
import xlsxwriter
from odoo import models, fields, api, _
from odoo.exceptions import ValidationError, UserError
import logging

_logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    openpyxl = None

class ExcelImportWizard(models.TransientModel):
    _name = 'cpk.reimburse.excel.import.wizard'
    _description = 'Excel 匯入憑證與請款明細精靈'

    reimburse_id = fields.Many2one('cpk.reimburse', '請款單', required=True)
    import_file = fields.Binary('Excel 檔案', required=True)
    import_filename = fields.Char('檔案名稱')
    
    # 匯入選項
    import_mode = fields.Selection([
        ('replace', '替換現有資料'),
        ('append', '附加到現有資料'),
    ], string='匯入模式', default='append', required=True)
    
    # 範例檔案
    template_file = fields.Binary('範例檔案', readonly=True)
    template_filename = fields.Char('範例檔案名稱', readonly=True)
    
    # 匯入結果
    import_log = fields.Text('匯入結果', readonly=True)
    voucher_count = fields.Integer('憑證數量', readonly=True)
    expense_line_count = fields.Integer('請款明細數量', readonly=True)

    @api.model
    def default_get(self, fields_list):
        """預設值設定"""
        res = super().default_get(fields_list)
        
        # 從 context 取得選取的請款單
        active_id = self.env.context.get('active_id')
        if active_id:
            res['reimburse_id'] = active_id
            
        return res

    def action_download_template(self):
        """下載範例檔案"""
        template_file, template_filename = self._generate_template_file()
        
        self.write({
            'template_file': template_file,
            'template_filename': template_filename,
        })
        
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def _generate_template_file(self):
        """產生範例 Excel 檔案"""
        if not xlsxwriter:
            raise UserError(_('請安裝 xlsxwriter 套件：pip install xlsxwriter'))

        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        
        # 定義樣式
        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#4CAF50',
            'font_color': 'white',
            'align': 'center',
            'valign': 'vcenter',
            'border': 1
        })
        
        required_format = workbook.add_format({
            'bold': True,
            'bg_color': '#FFC107',
            'font_color': 'black',
            'align': 'center',
            'valign': 'vcenter',
            'border': 1
        })
        
        example_format = workbook.add_format({
            'bg_color': '#E8F5E8',
            'border': 1
        })
        
        note_format = workbook.add_format({
            'font_color': 'red',
            'italic': True
        })

        # Sheet 1: 憑證資料
        voucher_sheet = workbook.add_worksheet('憑證資料')
        
        # 憑證表頭
        voucher_headers = [
            '憑證類型*', '發票日期', '發票號碼', '統一編號', 
            '稅碼', '交易幣別*', '含稅金額*', '未稅金額', '稅額'
        ]
        
        for col, header in enumerate(voucher_headers):
            if '*' in header:
                voucher_sheet.write(0, col, header, required_format)
            else:
                voucher_sheet.write(0, col, header, header_format)
        
        # 憑證範例資料
        voucher_examples = [
            ['invoice', '2024-01-15', 'AB12345678', '12345678', 'VA', 'TWD', 1050, 1000, 50],
            ['receipt', '2024-01-16', 'RC001', '', 'VN', 'TWD', 500, 500, 0],
        ]
        
        for row, example in enumerate(voucher_examples, 1):
            for col, value in enumerate(example):
                voucher_sheet.write(row, col, value, example_format)
        
        # 設定憑證欄位寬度
        voucher_sheet.set_column('A:A', 12)  # 憑證類型
        voucher_sheet.set_column('B:B', 12)  # 發票日期
        voucher_sheet.set_column('C:C', 15)  # 發票號碼
        voucher_sheet.set_column('D:D', 12)  # 統一編號
        voucher_sheet.set_column('E:E', 8)   # 稅碼
        voucher_sheet.set_column('F:F', 12)  # 交易幣別
        voucher_sheet.set_column('G:I', 12)  # 金額欄位
        
        # 憑證說明
        voucher_sheet.write('A5', '說明：', note_format)
        voucher_sheet.write('A6', '* 標記為必填欄位', note_format)
        voucher_sheet.write('A7', '憑證類型：invoice=發票, receipt=收據', note_format)
        voucher_sheet.write('A8', '日期格式：YYYY-MM-DD (例：2024-01-15)', note_format)
        voucher_sheet.write('A9', '金額計算：含稅金額 = 未稅金額 + 稅額', note_format)

        # Sheet 2: 請款明細
        expense_sheet = workbook.add_worksheet('請款明細')
        
        # 請款明細表頭
        expense_headers = [
            '憑證序號*', '費用類別代碼*', '成本中心', '專案代號', 
            '含稅金額*', '未稅金額', '稅額', '費用說明'
        ]
        
        for col, header in enumerate(expense_headers):
            if '*' in header:
                expense_sheet.write(0, col, header, required_format)
            else:
                expense_sheet.write(0, col, header, header_format)
        
        # 請款明細範例資料
        expense_examples = [
            [1, '0101', '資訊部', 'PRJ001', 525, 500, 25, '計程車車資'],
            [1, '0102', '資訊部', 'PRJ001', 525, 500, 25, '停車費'],
            [2, '0201', '業務部', '', 500, 500, 0, '餐費'],
        ]
        
        for row, example in enumerate(expense_examples, 1):
            for col, value in enumerate(example):
                expense_sheet.write(row, col, value, example_format)
        
        # 設定請款明細欄位寬度
        expense_sheet.set_column('A:A', 12)  # 憑證序號
        expense_sheet.set_column('B:B', 15)  # 費用類別代碼
        expense_sheet.set_column('C:C', 12)  # 成本中心
        expense_sheet.set_column('D:D', 12)  # 專案代號
        expense_sheet.set_column('E:G', 12)  # 金額欄位
        expense_sheet.set_column('H:H', 20)  # 費用說明
        
        # 請款明細說明
        expense_sheet.write('A6', '說明：', note_format)
        expense_sheet.write('A7', '* 標記為必填欄位', note_format)
        expense_sheet.write('A8', '憑證序號：對應憑證資料的行號（從1開始）', note_format)
        expense_sheet.write('A9', '費用類別代碼：請參考系統中的費用類別設定', note_format)
        expense_sheet.write('A10', '金額計算：含稅金額 = 未稅金額 + 稅額', note_format)

        workbook.close()
        output.seek(0)
        
        template_file = base64.b64encode(output.read())
        template_filename = f'憑證匯入範例_{fields.Date.today().strftime("%Y%m%d")}.xlsx'
        
        return template_file, template_filename

    def action_import_excel(self):
        """執行 Excel 匯入"""
        if not openpyxl:
            raise UserError(_('請安裝 openpyxl 套件：pip install openpyxl'))
        
        if not self.import_file:
            raise ValidationError(_('請選擇要匯入的 Excel 檔案'))
        
        try:
            # 解碼並讀取 Excel 檔案
            excel_data = base64.b64decode(self.import_file)
            workbook = openpyxl.load_workbook(io.BytesIO(excel_data), data_only=True)
            
            # 檢查工作表
            if '憑證資料' not in workbook.sheetnames:
                raise ValidationError(_('Excel 檔案中找不到「憑證資料」工作表'))
            if '請款明細' not in workbook.sheetnames:
                raise ValidationError(_('Excel 檔案中找不到「請款明細」工作表'))
            
            # 匯入前清理（如果選擇替換模式）
            if self.import_mode == 'replace':
                self.reimburse_id.voucher_ids.unlink()
                self.reimburse_id.expense_line_ids.unlink()
            
            # 匯入憑證資料
            voucher_sheet = workbook['憑證資料']
            vouchers = self._import_vouchers(voucher_sheet)
            
            # 匯入請款明細
            expense_sheet = workbook['請款明細']
            expense_lines = self._import_expense_lines(expense_sheet, vouchers)
            
            # 產生匯入報告
            self._generate_import_report(vouchers, expense_lines)
            
        except Exception as e:
            _logger.error(f'Excel 匯入錯誤: {str(e)}')
            raise UserError(_('匯入失敗：%s') % str(e))
        
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def _import_vouchers(self, sheet):
        """匯入憑證資料"""
        vouchers = []
        errors = []
        
        # 取得稅碼對照
        tax_codes = {tc.code: tc.id for tc in self.env['cpk.tax.code'].search([])}
        
        # 取得幣別對照
        currencies = {c.name: c.id for c in self.env['res.currency'].search([])}
        
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(row):  # 跳過空行
                continue
            
            try:
                voucher_type, invoice_date, invoice_number, tax_id, tax_code, \
                currency, amount_total, amount_untaxed, tax_amount = row[:9]
                
                # 驗證必填欄位
                if not voucher_type:
                    errors.append(f'行 {row_num}：憑證類型為必填')
                    continue
                
                if not currency:
                    errors.append(f'行 {row_num}：交易幣別為必填')
                    continue
                
                if not amount_total:
                    errors.append(f'行 {row_num}：含稅金額為必填')
                    continue
                
                # 驗證憑證類型
                if voucher_type not in ['invoice', 'receipt']:
                    errors.append(f'行 {row_num}：憑證類型必須是 invoice 或 receipt')
                    continue
                
                # 建立憑證資料
                voucher_vals = {
                    'reimburse_id': self.reimburse_id.id,
                    'voucher_type': voucher_type,
                    'amount_total': float(amount_total) if amount_total else 0,
                }
                
                # 選擇性欄位
                if invoice_date:
                    voucher_vals['invoice_date'] = invoice_date
                if invoice_number:
                    voucher_vals['invoice_number'] = str(invoice_number)
                if tax_id:
                    voucher_vals['tax_id'] = str(tax_id)
                if tax_code and tax_code in tax_codes:
                    voucher_vals['tax_code_id'] = tax_codes[tax_code]
                if currency in currencies:
                    voucher_vals['currency_id'] = currencies[currency]
                if amount_untaxed:
                    voucher_vals['amount_untaxed'] = float(amount_untaxed)
                if tax_amount:
                    voucher_vals['tax_amount'] = float(tax_amount)
                
                voucher = self.env['cpk.voucher'].create(voucher_vals)
                vouchers.append((row_num, voucher))
                
            except Exception as e:
                errors.append(f'行 {row_num}：{str(e)}')
                continue
        
        if errors:
            raise ValidationError(_('憑證資料匯入錯誤：\n%s') % '\n'.join(errors))
        
        return vouchers

    def _import_expense_lines(self, sheet, vouchers):
        """匯入請款明細"""
        expense_lines = []
        errors = []
        
        # 建立憑證索引對照
        voucher_map = {row_num: voucher for row_num, voucher in vouchers}
        
        # 取得費用類別對照
        expense_categories = {ec.code: ec.id for ec in self.env['cpk.expense.category'].search([])}
        
        # 取得部門對照
        departments = {dept.name: dept.id for dept in self.env['hr.department'].search([])}
        
        # 取得專案對照
        analytic_accounts = {aa.code: aa.id for aa in self.env['account.analytic.account'].search([]) if aa.code}
        
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(row):  # 跳過空行
                continue
            
            try:
                voucher_seq, category_code, cost_center, project_code, \
                amount_total, amount_untaxed, tax_amount, description = row[:8]
                
                # 驗證必填欄位
                if not voucher_seq:
                    errors.append(f'行 {row_num}：憑證序號為必填')
                    continue
                
                if not category_code:
                    errors.append(f'行 {row_num}：費用類別代碼為必填')
                    continue
                
                if not amount_total:
                    errors.append(f'行 {row_num}：含稅金額為必填')
                    continue
                
                # 驗證憑證序號
                voucher_row_num = int(voucher_seq) + 1  # Excel 行號轉換
                if voucher_row_num not in voucher_map:
                    errors.append(f'行 {row_num}：找不到憑證序號 {voucher_seq} 對應的憑證')
                    continue
                
                # 驗證費用類別
                if category_code not in expense_categories:
                    errors.append(f'行 {row_num}：找不到費用類別代碼 {category_code}')
                    continue
                
                # 建立請款明細
                expense_vals = {
                    'reimburse_id': self.reimburse_id.id,
                    'voucher_id': voucher_map[voucher_row_num].id,
                    'expense_category_id': expense_categories[category_code],
                    'amount_total': float(amount_total) if amount_total else 0,
                    'sequence': (len(expense_lines) + 1) * 10,
                }
                
                # 選擇性欄位
                if cost_center and cost_center in departments:
                    expense_vals['cost_center_id'] = departments[cost_center]
                if project_code and project_code in analytic_accounts:
                    expense_vals['analytic_account_id'] = analytic_accounts[project_code]
                if amount_untaxed:
                    expense_vals['amount_untaxed'] = float(amount_untaxed)
                if tax_amount:
                    expense_vals['tax_amount'] = float(tax_amount)
                if description:
                    expense_vals['description'] = str(description)
                
                expense_line = self.env['cpk.expense.line'].create(expense_vals)
                expense_lines.append(expense_line)
                
            except Exception as e:
                errors.append(f'行 {row_num}：{str(e)}')
                continue
        
        if errors:
            raise ValidationError(_('請款明細匯入錯誤：\n%s') % '\n'.join(errors))
        
        return expense_lines

    def _generate_import_report(self, vouchers, expense_lines):
        """產生匯入報告"""
        voucher_count = len(vouchers)
        expense_line_count = len(expense_lines)
        
        report_lines = [
            f'匯入完成！',
            f'',
            f'匯入統計：',
            f'- 憑證數量：{voucher_count} 筆',
            f'- 請款明細：{expense_line_count} 筆',
            f'',
            f'憑證明細：'
        ]
        
        for row_num, voucher in vouchers:
            report_lines.append(
                f'  行 {row_num-1}: {voucher.voucher_type} - '
                f'{voucher.amount_total} {voucher.currency_id.name}'
            )
        
        if expense_lines:
            report_lines.append(f'')
            report_lines.append(f'請款明細：')
            for expense_line in expense_lines:
                report_lines.append(
                    f'  {expense_line.expense_category_id.name}: '
                    f'{expense_line.amount_total} {self.reimburse_id.currency_id.name}'
                )
        
        self.write({
            'import_log': '\n'.join(report_lines),
            'voucher_count': voucher_count,
            'expense_line_count': expense_line_count,
        })

    def action_close(self):
        """關閉精靈並重新整理父視窗"""
        return {
            'type': 'ir.actions.act_window_close',
        }